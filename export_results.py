"""导出投票数据备份（可对线上 Render 数据库或本地库运行）。

用途：Render 免费 Postgres 自创建起 30 天过期、过期后 14 天被彻底删除，
且不支持任何形式的备份。投票结束后务必用本脚本把数据导出留存。

用法：
    # 导出线上数据（把 <外部连接串> 换成 Render 控制台 Info 页的 External Database URL）
    DATABASE_URL="<外部连接串>" python export_results.py

    # 导出本地 SQLite 数据
    python export_results.py

产物（默认写到 exports/ 目录）：
    vote_export_<时间戳>.xlsx   多工作表汇总，可直接用 Excel 打开
    vote_export_<时间戳>.json   完整结构化备份，便于以后重新导入
"""
import json
import os
import sys
from datetime import datetime

# 让脚本可独立运行：复用项目自身的 app/models 配置
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app import app  # noqa: E402  （导入即完成配置与建表）
from models import Ballot, Poll, PollOption, User, Voter, db  # noqa: E402

OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "exports")


def collect():
    """把所有投票及其名单、选项、选票整理成可序列化结构。"""
    data = {
        "exported_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "database": (app.config.get("SQLALCHEMY_DATABASE_URI") or "").split("@")[-1],
        "polls": [],
    }

    for poll in Poll.query.order_by(Poll.created_at).all():
        options = PollOption.query.filter_by(poll_id=poll.id).order_by(
            PollOption.position).all()
        voters = Voter.query.filter_by(poll_id=poll.id).order_by(Voter.name).all()
        ballots = Ballot.query.filter_by(poll_id=poll.id).all()

        opt_name = {o.id: o.name for o in options}
        voter_name = {v.id: v.name for v in voters}

        counts = {o.id: 0 for o in options}
        for b in ballots:
            counts[b.option_id] = counts.get(b.option_id, 0) + 1
        total = sum(counts.values())

        creator = db.session.get(User, poll.creator_id)
        data["polls"].append({
            "id": poll.id,
            "title": poll.title,
            "description": poll.description or "",
            "creator": creator.real_name if creator else "?",
            "created_at": poll.created_at.strftime("%Y-%m-%d %H:%M:%S") if poll.created_at else None,
            "closed_at": poll.closed_at.strftime("%Y-%m-%d %H:%M:%S") if poll.closed_at else None,
            "is_closed": bool(poll.is_closed),
            "is_published": bool(poll.is_published),
            "allow_self_vote": bool(poll.allow_self_vote),
            "allow_abstain": bool(poll.allow_abstain),
            "total_voters": len(voters),
            "total_votes": total,
            "results": [{
                "option": o.name,
                "is_abstain": bool(o.is_abstain),
                "count": counts.get(o.id, 0),
                "percent": round(counts.get(o.id, 0) * 100 / total, 1) if total else 0.0,
            } for o in options],
            "voters": [{
                "name": v.name,
                "registered": v.user_id is not None,
                "has_voted": bool(v.has_voted),
                "voted_at": v.voted_at.strftime("%Y-%m-%d %H:%M:%S") if v.voted_at else None,
            } for v in voters],
            # 选票明细：发起人可见的「谁投了谁」，这是最关键、最无法重建的数据
            "ballots": sorted([{
                "voter": voter_name.get(b.voter_id, "?"),
                "option": opt_name.get(b.option_id, "?"),
                "created_at": b.created_at.strftime("%Y-%m-%d %H:%M:%S") if b.created_at else None,
            } for b in ballots], key=lambda x: x["voter"]),
        })
    return data


def write_json(data, stamp):
    path = os.path.join(OUT_DIR, "vote_export_%s.json" % stamp)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return path


def write_xlsx(data, stamp):
    """写多工作表 Excel。openpyxl 已在 requirements.txt 中。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        print("  ! 未安装 openpyxl，跳过 Excel 导出（JSON 已生成）")
        return None

    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "投票汇总"
    ws.append(["投票ID", "标题", "发起人", "创建时间", "结束时间",
               "状态", "应投人数", "已投票数", "投票率"])
    for c in ws[1]:
        c.font = bold
    for p in data["polls"]:
        rate = "%.1f%%" % (p["total_votes"] * 100 / p["total_voters"]) if p["total_voters"] else "-"
        ws.append([p["id"], p["title"], p["creator"], p["created_at"], p["closed_at"],
                   "已结束" if p["is_closed"] else "进行中",
                   p["total_voters"], p["total_votes"], rate])

    ws2 = wb.create_sheet("票数结果")
    ws2.append(["投票ID", "标题", "选项", "是否弃票", "票数", "占比"])
    for c in ws2[1]:
        c.font = bold
    for p in data["polls"]:
        for r in p["results"]:
            ws2.append([p["id"], p["title"], r["option"],
                        "是" if r["is_abstain"] else "", r["count"], "%.1f%%" % r["percent"]])

    ws3 = wb.create_sheet("名单与投票状态")
    ws3.append(["投票ID", "标题", "姓名", "是否注册账号", "是否已投", "投票时间"])
    for c in ws3[1]:
        c.font = bold
    for p in data["polls"]:
        for v in p["voters"]:
            ws3.append([p["id"], p["title"], v["name"],
                        "是" if v["registered"] else "否",
                        "是" if v["has_voted"] else "否", v["voted_at"]])

    ws4 = wb.create_sheet("选票明细")
    ws4.append(["投票ID", "标题", "投票人", "所投选项", "投票时间"])
    for c in ws4[1]:
        c.font = bold
    for p in data["polls"]:
        for b in p["ballots"]:
            ws4.append([p["id"], p["title"], b["voter"], b["option"], b["created_at"]])

    for sheet in wb.worksheets:
        for col in sheet.columns:
            width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(width + 4, 10), 45)

    path = os.path.join(OUT_DIR, "vote_export_%s.xlsx" % stamp)
    wb.save(path)
    return path


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    with app.app_context():
        data = collect()

    print("数据库：%s" % data["database"])
    print("投票场次：%d" % len(data["polls"]))
    for p in data["polls"]:
        print("  [%d] %s — 应投 %d 人，已投 %d 票%s"
              % (p["id"], p["title"], p["total_voters"], p["total_votes"],
                 "（已结束）" if p["is_closed"] else ""))

    if not data["polls"]:
        print("\n! 数据库中没有任何投票记录，请确认 DATABASE_URL 指向正确的库。")

    jp = write_json(data, stamp)
    xp = write_xlsx(data, stamp)
    print("\n已导出：")
    print("  %s" % jp)
    if xp:
        print("  %s" % xp)


if __name__ == "__main__":
    main()
